#!/usr/bin/env python3
"""Instagram audience analysis from CSV follower exports.

Works with files from the Follower Tracker extension and from other services: columns
(name, link or username, status) are detected automatically by header and content.
Python standard library only; sends nothing to the network.

Commands:
  inspect FILE                 columns, row count, how name, link and status were detected
  changes FILE                 status table (new / unfollowed / stayed):
                               numbers per group, CSV per group in --out-dir
  compare OLD NEW              two lists from different dates: who arrived and who left
                               (--out-gained/--out-lost FILE to save as CSV)
  mutual FOLLOWERS FOLLOWING   mutual follows, not following back, fans
  profile FILE                 approximate audience composition from names and usernames
  sample FILE                  random sample of profiles for manual review

Global options (go BEFORE the command):
  --name-col NAME   profile name column
  --url-col NAME    link (or username) column
  --status-col NAME status column
  --json            JSON output
changes options: --out-dir DIR --map "value=new|unfollowed|stayed" (repeatable) --show 30
sample options: --n 40 --seed 1 --out sample.md
profile options: --examples 5
suspects FILE --out OUT.xlsx|OUT.csv --lang en|ru
                             suspicious accounts (by confidence) + a separate "possibly real" list;
                             .xlsx needs openpyxl (falls back to two CSV files if it's missing)
"""
import argparse, csv, io, json, random, re, sys, unicodedata
from collections import Counter, defaultdict

URL_HINTS = ["url", "link", "href", "profile", "ссылка", "профиль", "username", "user name", "login", "логин",
             "handle", "nick", "ник", "account", "аккаунт", "instagram", "инстаграм"]
HANDLE_RE = re.compile(r"^@?[A-Za-z0-9._]{1,30}$")
NAME_HINTS = ["full name", "fullname", "full_name", "name", "имя", "display"]

# Keywords in a profile name -> field. Rough heuristic, indicative only.
# Keywords cover English and Russian, since profile names come in both.
SPHERES = {
    "Beauty": ["nail", "ногт", "маникюр", "manicure", "brow", "бров", "lash", "ресниц", "визаж", "makeup", "make up",
                "beauty", "бьюти", "косметолог", "cosmetolog", "hair", "волос", "парикмах", "barber", "барбер", "salon", "салон", "spa", "спа", "эпиляц", "депиляц"],
    "Health and psychology": ["психолог", "psycholog", "психотерап", "therap", "нутрициолог", "nutrition", "врач", "doctor", "dr.",
                              "стоматолог", "dent", "клиник", "clinic", "массаж", "massage", "остеопат", "health", "здоров"],
    "Sports and fitness": ["fitness", "фитнес", "trainer", "тренер", "coach fit", "yoga", "йога", "pilates", "пилатес", "sport", "спорт", "gym"],
    "Food and restaurants": ["cafe", "кафе", "кофе", "coffee", "restaurant", "ресторан", "food", "еда", "кондитер", "торт", "cake",
                        "bakery", "пекарн", "chef", "шеф", "pizza", "пицц", "sushi", "суши", "доставка еды", "кулинар"],
    "Education": ["school", "школ", "курс", "course", "academy", "академ", "обучени", "учител", "teacher", "репетитор", "tutor", "english", "английск", "university", "университет"],
    "Marketing and SMM": ["smm", "маркет", "market", "таргет", "target", "агентств", "agency", "продвижен", "promo", "контент", "content", "копирайт", "copywrit", "seo"],
    "Photo, video, design": ["photo", "фото", "video", "видео", "reels", "дизайн", "design", "illustr", "иллюстр", "художник", "artist", "creative", "креатив"],
    "Fashion and shops": ["одежд", "wear", "fashion", "мода", "стилист", "stylist", "boutique", "бутик", "showroom", "шоурум", "jewel", "украшен", "handmade", "хендмейд"],
    "Real estate, renovation, construction": ["недвиж", "realty", "estate", "риелтор", "realtor", "ремонт", "repair", "строит", "build", "интерьер", "interior", "мебел", "furniture"],
    "Travel": ["travel", "путешеств", "tour", "отель", "hotel", "hostel"],
    "Business and finance": ["бизнес", "business", "предприним", "entrepreneur", "финанс", "finance", "инвест", "invest", "юрист", "lawyer", "адвокат", "бухгалт", "account", "консалт", "consult", "ceo", "founder", "основател", "эксперт", "expert"],
    "Bloggers and media": ["blog", "блог", "influenc", "media", "медиа", "podcast", "подкаст", "news", "новост", "журнал", "magazine"],
}
BIZ_MARKERS = ["official", "официальн", "shop", "store", "studio", "студия", "салон", "магазин", "доставка", "заказ", "order", "ltd", "llc", "ооо", "ип "]
SEPARATORS = re.compile(r"[|•·/✨🔥💅📍—–]|\s-\s")
VOWELS = set("aeiouyаеёиоуыэюя")

STATUS_HINTS = ["status", "статус", "state", "состояние", "change", "изменени", "event", "событ"]
# Status values from Follower Tracker and common variants from other services (English and Russian).
STATUS_VALUES = {
    "new": ["recent follow", "recent follower", "new", "new follower", "followed", "follow", "follower gained",
            "gained", "joined", "subscribed", "новый", "новая", "новые", "новый подписчик", "подписался",
            "подписалась", "подписались", "пришёл", "пришел", "пришла"],
    "unfollowed": ["unfollowed", "unfollow", "unfollower", "lost", "left", "removed", "gone", "unsubscribed",
                   "отписался", "отписалась", "отписались", "отписка", "ушёл", "ушел", "ушла", "потерян"],
    "stayed": ["stayed", "stay", "kept", "remained", "remain", "existing", "same", "unchanged", "old",
               "остался", "осталась", "остались", "без изменений", "прежний"],
}
GROUP_LABELS = {"new": "new (recent follow)", "unfollowed": "unfollowed", "stayed": "stayed"}


def read_csv(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = raw.decode(enc); break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    rows = [r for r in reader if any((v or "").strip() for v in r.values() if isinstance(v, str))]
    return reader.fieldnames or [], rows


def norm_status(v):
    return re.sub(r"[\s_\-]+", " ", (v or "").strip().lower())


def status_lookup(extra=None):
    table = {norm_status(v): g for g, vals in STATUS_VALUES.items() for v in vals}
    for item in extra or []:
        if "=" not in item:
            sys.exit(f"--map expects 'value=new|unfollowed|stayed', got: {item}")
        val, grp = item.rsplit("=", 1)
        if grp.strip() not in STATUS_VALUES:
            sys.exit(f"Group must be new, unfollowed or stayed, got: {grp}")
        table[norm_status(val)] = grp.strip()
    return table


def detect_status_col(fields, rows, forced=None):
    if forced:
        if forced not in fields:
            sys.exit(f"Column '{forced}' not found. Available: {fields}")
        return forced
    low = {f.lower().strip(): f for f in fields}
    by_name = next((orig for h in STATUS_HINTS for l, orig in low.items() if h in l), None)
    if by_name:
        return by_name
    known = status_lookup()
    sample = rows[:300]
    for f in fields:  # by content: nearly all values are known statuses
        vals = [norm_status(r.get(f)) for r in sample if (r.get(f) or "").strip()]
        if vals and sum(v in known for v in vals) / len(vals) >= 0.8:
            return f
    return None


def detect_cols(fields, rows, name_col=None, url_col=None, status_col=None):
    fields = [f for f in fields if f != status_col]
    for forced in (name_col, url_col):
        if forced and forced not in fields:
            sys.exit(f"Column '{forced}' not found. Available: {fields}")
    if not url_col:
        # by content: the column where instagram.com appears most
        best, score = None, 0
        for f in fields:
            s = sum(1 for r in rows[:200] if "instagram.com" in (r.get(f) or "").lower())
            if s > score:
                best, score = f, s
        if not best:
            low = {f.lower().strip(): f for f in fields}
            best = next((orig for h in URL_HINTS for l, orig in low.items() if h in l), None)
        if not best:  # by content: the column where nearly all values look like usernames
            for f in fields:
                vals = [(r.get(f) or "").strip() for r in rows[:200] if (r.get(f) or "").strip()]
                if vals and sum(bool(HANDLE_RE.match(v)) for v in vals) / len(vals) >= 0.9 \
                        and len(set(vals)) / len(vals) > 0.9:
                    best = f; break
        url_col = best
    if not name_col:
        low = {f.lower().strip(): f for f in fields if f != url_col}
        name_col = next((orig for h in NAME_HINTS for l, orig in low.items() if h in l), None)
        if not name_col:
            rest = [f for f in fields if f != url_col]
            name_col = rest[0] if rest else None
    return name_col, url_col


def username_of(value):
    v = (value or "").strip()
    m = re.search(r"instagram\.com/([^/?#\s]+)", v, re.I)
    if m:
        v = m.group(1)
    return v.lstrip("@").strip("/").lower()


def load(path, a):
    fields, rows = read_csv(path)
    status_col = detect_status_col(fields, rows, a.status_col)
    name_col, url_col = detect_cols(fields, rows, a.name_col, a.url_col, status_col)
    a._status_col = status_col
    people = {}
    for r in rows:
        key = username_of(r.get(url_col)) if url_col else ""
        name = (r.get(name_col) or "").strip() if name_col else ""
        if not key:
            key = name.lower()
        if key:
            people[key] = {"username": key, "name": name, "url": (r.get(url_col) or "").strip() if url_col else "",
                           "status": (r.get(status_col) or "").strip() if status_col else ""}
    return fields, rows, name_col, url_col, people


def out(result, as_json):
    if as_json:
        print(json.dumps(result, ensure_ascii=False, indent=2)); return
    for k, v in result.items():
        if isinstance(v, list):
            print(f"\n{k} ({len(v)}):")
            for x in v[:200]:
                print(f"  {x}")
            if len(v) > 200:
                print(f"  ... and {len(v) - 200} more")
        elif isinstance(v, dict):
            print(f"\n{k}:")
            for kk, vv in v.items():
                print(f"  {kk}: {vv}")
        else:
            print(f"{k}: {v}")


def pct(n, total):
    return round(n / total * 100, 1) if total else 0.0


def cmd_inspect(a):
    fields, rows, nc, uc, people = load(a.file, a)
    sc = a._status_col
    res = {"file": a.file, "rows": len(rows), "unique profiles": len(people),
           "columns": ", ".join(fields), "name column": nc, "link column": uc,
           "status column": sc or "none (plain follower list)"}
    if sc:
        known = status_lookup()
        vals = Counter(p["status"] for p in people.values())
        res["status values"] = {(v or "<empty>"): f"{c} → {GROUP_LABELS.get(known.get(norm_status(v)), 'NOT RECOGNIZED')}"
                                   for v, c in vals.most_common()}
    res["sample rows"] = [json.dumps(p, ensure_ascii=False) for p in list(people.values())[:3]]
    out(res, a.json)


def write_list(path, items):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["Name", "Profile URL"])
        for p in sorted(items, key=lambda x: x["username"]):
            w.writerow([p["name"], p["url"] or f"https://www.instagram.com/{p['username']}/"])


def cmd_changes(a):
    import os
    *_, people = load(a.file, a)
    if not a._status_col:
        sys.exit("No status column found in the file. Set it with --status-col, "
                 "or compare two lists from different dates with the compare command.")
    table = status_lookup(a.map)
    groups, unknown = defaultdict(list), Counter()
    for p in people.values():
        g = table.get(norm_status(p["status"]))
        if g:
            groups[g].append(p)
        else:
            unknown[p["status"] or "<empty>"] += 1
    new, lost, stayed = len(groups["new"]), len(groups["unfollowed"]), len(groups["stayed"])
    before, after = stayed + lost, stayed + new
    res = {"status column": a._status_col,
           "before (stayed + unfollowed)": before, "after (stayed + new)": after,
           "arrived": new, "left": lost, "net change": new - lost,
           "growth, % of previous base": pct(new - lost, before), "churn, % of previous base": pct(lost, before)}
    if not stayed:
        res["warning"] = ("The file has no 'stayed' rows, so the previous base can't be calculated "
                          "and percentages aren't meaningful. Use the absolute numbers.")
    if unknown:
        res["UNRECOGNIZED status values (ask the user, then pass --map)"] = dict(unknown.most_common())
    if a.out_dir:
        os.makedirs(a.out_dir, exist_ok=True)
        for g in ("new", "unfollowed", "stayed"):
            if groups[g]:
                write_list(os.path.join(a.out_dir, f"{g}.csv"), groups[g])
        res["saved"] = f"{a.out_dir}/new.csv, unfollowed.csv, stayed.csv; run profile on each"
    for g in ("new", "unfollowed"):
        names = sorted(p["username"] for p in groups[g])
        res[f"{GROUP_LABELS[g]}, first {a.show}"] = names[:a.show]
    out(res, a.json)


def cmd_compare(a):
    *_, old = load(a.old, a)
    *_, new = load(a.new, a)
    o, n = set(old), set(new)
    gained, lost = n - o, o - n
    for path, keys, src in ((a.out_gained, gained, new), (a.out_lost, lost, old)):
        if path:
            write_list(path, [src[k] for k in keys])
    base = len(o) or 1
    out({"before": len(o), "after": len(n), "net change": len(n) - len(o),
         "growth, %": round((len(n) - len(o)) / base * 100, 2),
         "arrived": len(gained), "left": len(lost), "churn, % of base": round(len(lost) / base * 100, 2),
         "new followers": sorted(gained), "unfollowers": sorted(lost)}, a.json)


def cmd_mutual(a):
    *_, fol = load(a.followers, a)
    *_, fng = load(a.following, a)
    f, g = set(fol), set(fng)
    out({"followers": len(f), "following": len(g), "mutual": sorted(f & g),
         "not following back": sorted(g - f), "fans (you don't follow them)": sorted(f - g)}, a.json)


def script_of(name):
    cyr = len(re.findall(r"[а-яёА-ЯЁ]", name)); lat = len(re.findall(r"[a-zA-Z]", name))
    other = len(re.findall(r"[^\W\d_a-zA-Zа-яёА-ЯЁ]", name))
    if not (cyr or lat or other):
        return "no letters / empty"
    top = max((cyr, "Cyrillic"), (lat, "Latin"), (other, "other scripts"))
    return top[1]


DATE_RE = re.compile(
    r"(19[4-9]\d|20[0-2]\d)"                              # a birth or other year: 1940-2029
    r"|(0[1-9]|[12]\d|3[01])[.\-_]?(0[1-9]|1[0-2])[.\-_]?(\d{2}|\d{4})(?!\d)"  # ddmmyy / dd.mm.yyyy
    r"|(?<!\d)\d{2}[.\-_]?(0[1-9]|1[0-2])[.\-_]?(0[1-9]|[12]\d|3[01])(?!\d)"   # yymmdd / yy.mm.dd
)


def _random_letters(text):
    text = unicodedata.normalize("NFKC", text)  # decorative fonts (𝒜𝓃𝒶) -> plain letters
    for part in re.split(r"[\W._\d]+", text.lower()):
        lt = [c for c in part if c.isalpha()]
        if any(not re.match(r"[a-zа-яё]", c) for c in lt):
            continue  # vowel check only works for Latin and Cyrillic; skip other scripts
        run = max((len(m) for m in re.findall(r"[^aeiouyаеёиоуыэюя\W\d_]+", part)), default=0)
        if len(lt) >= 7 and sum(c in VOWELS for c in lt) / len(lt) < 0.2 and run >= 5:
            return True
    return False


def suspicious(username, name):
    """Signs of a bot or throwaway account, with a rough score.

    Returns (reasons, score). score 0 = not flagged. These are "check this" signals,
    not proof: real people also use digits in usernames (often a birth date) or leave the name empty.
    """
    reasons, score = [], 0
    u = re.sub(r"[._]", "", username)
    digits = sum(c.isdigit() for c in u)
    ratio = digits / max(len(u), 1)
    name_letters = [c for c in name if c.isalpha()]
    has_name = len(name_letters) >= 2 and not _random_letters(name)
    date_like = bool(DATE_RE.search(username))

    if _random_letters(username):
        reasons.append("random-looking letters in username"); score += 2
    if name_letters and _random_letters(name):
        reasons.append("random-looking letters in name"); score += 1
    many_digits = digits >= 6 or (digits >= 4 and ratio > 0.5)
    username_letters = sum(c.isalpha() for c in u)
    if not name_letters and digits >= 3:
        if date_like:
            # a name plus a birth date with an empty profile name is common for real people
            reasons.append("no name; digits in username look like a date"); score += 1
            if username_letters == 0:
                reasons.append("username has no letters"); score += 1
        else:
            reasons.append("no name + digits in username"); score += 2
            if many_digits:
                reasons.append("username is mostly digits" if ratio > 0.6 else "many digits in username"); score += 1
    elif many_digits and not date_like:
        # digits next to a real name: weak signal; digits that look like a date: not flagged
        reasons.append("many digits in username"); score += 1
    return reasons, score


def confidence(score):
    return "high" if score >= 3 else "medium" if score == 2 else "low"


def possibly_real(username, name):
    """Accounts NOT flagged as suspicious only because the many digits in the username look like a
    birth date and the profile has a real name. Usually real people; listed separately for an optional check."""
    u = re.sub(r"[._]", "", username)
    digits = sum(c.isdigit() for c in u)
    many_digits = digits >= 6 or (digits >= 4 and digits / max(len(u), 1) > 0.5)
    name_letters = [c for c in name if c.isalpha()]
    has_name = len(name_letters) >= 2 and not _random_letters(name)
    return many_digits and has_name and bool(DATE_RE.search(username)) and not suspicious(username, name)[1]


_SPHERE_RE = {s: re.compile(r"(?<![^\W_])(?:" + "|".join(re.escape(k) for k in kws) + ")", re.I)
              for s, kws in SPHERES.items()}


def classify(text):
    # a keyword must start at a word boundary, so a short keyword isn't matched inside a longer word
    return [s for s, rx in _SPHERE_RE.items() if rx.search(text)]


def cmd_profile(a):
    *_, people = load(a.file, a)
    total = len(people)
    sphere_count, sphere_ex = Counter(), defaultdict(list)
    scripts, kinds = Counter(), Counter()
    susp, susp_rows, unclassified_desc = [], [], []
    for p in people.values():
        name, user = p["name"], p["username"]
        scripts[script_of(name)] += 1
        spheres = classify(name + " " + user.replace("_", " ").replace(".", " "))
        has_desc = bool(SEPARATORS.search(name)) or len(name.split()) >= 4
        biz = any(m in (" " + name.lower() + " " + user) for m in BIZ_MARKERS)
        if not name.strip():
            kinds["empty name"] += 1
        elif spheres or biz:
            kinds["looks like a business/specialist (field in name)"] += 1
        elif has_desc:
            kinds["name has a description, field not recognized"] += 1
            unclassified_desc.append(f"{user} — {name}")
        else:
            kinds["name only (probably a personal profile)"] += 1
        for s in spheres:
            sphere_count[s] += 1
            if len(sphere_ex[s]) < a.examples and (name or user) not in sphere_ex[s]:
                sphere_ex[s].append(name or user)
        r, score = suspicious(user, name)
        if score:
            susp.append(f"{user} [{confidence(score)}] ({', '.join(r)})")
            susp_rows.append((p, r, score))
    maybe_real = sum(possibly_real(p["username"], p["name"]) for p in people.values())
    out({
        "profiles": total,
        "profile type by name": {k: f"{v} ({pct(v, total)}%)" for k, v in kinds.most_common()},
        "fields by keyword (a profile can be in several)": {
            s: f"{c} ({pct(c, total)}%) — e.g.: {'; '.join(sphere_ex[s])}" for s, c in sphere_count.most_common()},
        "name alphabet": {k: f"{v} ({pct(v, total)}%)" for k, v in scripts.most_common()},
        "suspicious — check by hand": f"{len(susp)} ({pct(len(susp), total)}%): " + ", ".join(
            f"{lvl} {sum(confidence(x[2]) == lvl for x in susp_rows)}" for lvl in ("high", "medium", "low")),
        "suspicious examples": susp[:30],
        "possibly real (birth date in username + real name), not counted above": maybe_real,
        "to save these lists to a file": "run the suspects command",
        "described but field not recognized (classify these yourself)": unclassified_desc[:300],
        "note": "Heuristic based on names and usernames: indicative, not a precise portrait. Most people put only their name in the profile name.",
    }, a.json)


SUSPECT_TEXT = {
    "en": {
        "levels": {"high": "high", "medium": "medium", "low": "low"},
        "reasons": {},
        "possibly_reason": "birth date in username, real name",
        "head": ["Confidence", "Username", "Name", "Profile URL", "Why flagged", "Checked (bot / real)"],
        "check_values": "bot,real,not sure",
        "sheet_guide": "How to use", "sheet_susp": "Suspicious", "sheet_maybe": "Possibly real",
        "title": "Suspicious followers",
        "source": "Source file: {file}. Profiles in the file: {total}.",
        "summary": "Summary", "total": "Flagged as suspicious", "by_level": "— {lvl} confidence",
        "maybe_total": "Possibly real (separate tab)", "checked_bot": "Checked: bot", "checked_real": "Checked: real",
        "how_read": "How to read it",
        "read": [
            "High — several signs at once, e.g. no name and a username of mostly digits. Start checking with these.",
            "Medium — one strong sign, e.g. no name and digits in the username, or random-looking letters.",
            "Low — a weak sign, e.g. many digits in the username with a filled-in name. Many of these are real people.",
            "Possibly real (separate tab) — a birth date in the username and a real name. Usually real people; check only if you want a thorough cleanup.",
        ],
        "how_check": "How to check",
        "check": [
            "1. Open the profile from the link on the Suspicious tab.",
            "2. In the Checked column (yellow cells), pick from the list: bot, real or not sure.",
            "3. Remove bots by hand and gradually — a few dozen a day. Mass removal can get the account restricted by Instagram.",
        ],
        "disclaimer": "The signs are signals, not proof: real people also have digits in their usernames or an empty name.",
    },
    "ru": {
        "levels": {"high": "высокая", "medium": "средняя", "low": "низкая"},
        "reasons": {
            "random-looking letters in username": "случайный набор букв в нике",
            "random-looking letters in name": "случайный набор букв в имени",
            "no name + digits in username": "нет имени + цифры в нике",
            "no name; digits in username look like a date": "нет имени; цифры в нике похожи на дату",
            "username has no letters": "в нике нет букв",
            "username is mostly digits": "ник почти из одних цифр",
            "many digits in username": "много цифр в нике",
        },
        "possibly_reason": "дата рождения в нике, заполнено имя",
        "head": ["Уверенность", "Ник", "Имя профиля", "Ссылка на профиль", "Почему отмечен", "Проверено (бот / живой)"],
        "check_values": "бот,живой,не понятно",
        "sheet_guide": "Как пользоваться", "sheet_susp": "Подозрительные", "sheet_maybe": "Скорее живые",
        "title": "Подозрительные подписчики",
        "source": "Файл: {file}. Профилей в файле: {total}.",
        "summary": "Сводка", "total": "Подозрительных всего", "by_level": "— {lvl} уверенность",
        "maybe_total": "Скорее живые (отдельная вкладка)", "checked_bot": "Проверено: бот", "checked_real": "Проверено: живой",
        "how_read": "Как читать",
        "read": [
            "Высокая — несколько признаков сразу, например нет имени и ник почти из одних цифр. Начните проверку с них.",
            "Средняя — один сильный признак, например нет имени и цифры в нике, или случайный набор букв.",
            "Низкая — слабый признак, например много цифр в нике при заполненном имени. Многие из них живые люди.",
            "Скорее живые (отдельная вкладка) — дата рождения в нике и заполненное имя. Обычно это живые люди; проверяйте, только если хотите тщательную чистку.",
        ],
        "how_check": "Как проверять",
        "check": [
            "1. Откройте профиль по ссылке во вкладке «Подозрительные».",
            "2. В колонке «Проверено» (жёлтые ячейки) выберите из списка: бот, живой или не понятно.",
            "3. Удаляйте ботов вручную и постепенно — несколько десятков в день. Массовое удаление может привести к ограничениям аккаунта в Instagram.",
        ],
        "disclaimer": "Признаки — это сигналы, а не доказательство: у живых людей тоже бывают ники с цифрами или пустое имя.",
    },
}


def cmd_suspects(a):
    import os
    *_, people = load(a.file, a)
    T = SUSPECT_TEXT[a.lang]
    link = lambda p: p["url"] if p["url"].startswith("http") else f"https://www.instagram.com/{p['username']}/"
    susp, maybe = [], []
    for p in people.values():
        r, sc = suspicious(p["username"], p["name"])
        if sc:
            susp.append((sc, p, "; ".join(T["reasons"].get(x, x) for x in r)))
        elif possibly_real(p["username"], p["name"]):
            maybe.append((0, p, T["possibly_reason"]))
    susp.sort(key=lambda x: (-x[0], x[1]["username"])); maybe.sort(key=lambda x: x[1]["username"])
    rows = lambda items, lvl=None: [[lvl or T["levels"][confidence(sc)], p["username"], p["name"], link(p), why, ""]
                                    for sc, p, why in items]
    susp_rows, maybe_rows = rows(susp), rows(maybe, "—")

    xlsx = a.out.lower().endswith(".xlsx")
    if xlsx:
        try:
            import openpyxl  # optional: only used when it's already installed
        except ImportError:
            xlsx = False
            a.out = a.out[:-5] + ".csv"
            print("openpyxl is not installed: saving CSV files instead of .xlsx")
    if not xlsx:
        base = a.out[:-4] if a.out.lower().endswith(".csv") else a.out
        for path, data in ((base + ".csv", susp_rows), (base + "-possibly-real.csv", maybe_rows)):
            with open(path, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f); w.writerow(T["head"]); w.writerows(data)
        saved = f"{base}.csv, {base}-possibly-real.csv"
    else:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.worksheet.datavalidation import DataValidation
        FONT = "Arial"
        fills = {T["levels"]["high"]: "F4C7C3", T["levels"]["medium"]: "FCE8B2", T["levels"]["low"]: "E8EAED", "—": "E6F4EA"}
        wb = Workbook()
        guide = wb.active; guide.title = T["sheet_guide"]

        def data_sheet(title, data):
            ws = wb.create_sheet(title); ws.append(T["head"])
            for c in ws[1]:
                c.font = Font(name=FONT, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="3C4043")
                c.alignment = Alignment(vertical="center", wrap_text=True)
            for row in data:
                ws.append(row); i = ws.max_row
                for c in ws[i]: c.font = Font(name=FONT)
                ws.cell(i, 1).fill = PatternFill("solid", fgColor=fills.get(row[0], "FFFFFF"))
                ws.cell(i, 4).hyperlink = row[3]; ws.cell(i, 4).font = Font(name=FONT, color="1A73E8", underline="single")
                ws.cell(i, 6).fill = PatternFill("solid", fgColor="FFF8C5")
            for col, wdt in zip("ABCDEF", [13, 26, 30, 44, 46, 24]): ws.column_dimensions[col].width = wdt
            ws.freeze_panes = "A2"; last = max(ws.max_row, 2)
            ws.auto_filter.ref = f"A1:F{last}"
            dv = DataValidation(type="list", formula1=f'"{T["check_values"]}"', allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"F2:F{last}")
            return ws, last

        _, n1 = data_sheet(T["sheet_susp"], susp_rows)
        _, n2 = data_sheet(T["sheet_maybe"], maybe_rows)
        sq = lambda t: "'" + t.replace("'", "''") + "'"
        S1, S2 = sq(T["sheet_susp"]), sq(T["sheet_maybe"])
        bot, real = T["check_values"].split(",")[:2]
        lines = [(T["title"], None), (T["source"].format(file=os.path.basename(a.file), total=len(people)), None), ("", None),
                 (T["summary"], None), (T["total"], f"=COUNTA({S1}!A2:A{n1})")]
        lines += [(T["by_level"].format(lvl=T["levels"][k]), f'=COUNTIF({S1}!A2:A{n1},"{T["levels"][k]}")')
                  for k in ("high", "medium", "low")]
        lines += [(T["maybe_total"], f"=COUNTA({S2}!A2:A{n2})"),
                  (T["checked_bot"], f'=COUNTIF({S1}!F2:F{n1},"{bot}")+COUNTIF({S2}!F2:F{n2},"{bot}")'),
                  (T["checked_real"], f'=COUNTIF({S1}!F2:F{n1},"{real}")+COUNTIF({S2}!F2:F{n2},"{real}")'),
                  ("", None), (T["how_read"], None)] + [(x, None) for x in T["read"]]
        lines += [("", None), (T["how_check"], None)] + [(x, None) for x in T["check"]] + [("", None), (T["disclaimer"], None)]
        heads = {T["summary"], T["how_read"], T["how_check"]}
        for text, formula in lines:
            guide.append([text, formula]); i = guide.max_row
            guide.cell(i, 1).font = Font(name=FONT, bold=text in heads)
            if formula: guide.cell(i, 2).font = Font(name=FONT, bold=True)
        guide["A1"].font = Font(name=FONT, bold=True, size=14)
        guide.column_dimensions["A"].width = 110; guide.column_dimensions["B"].width = 12
        wb.active = 0; wb.save(a.out); saved = a.out
    by = Counter(confidence(sc) for sc, _, _ in susp)
    out({"profiles": len(people), "suspicious": len(susp),
         "by confidence": {k: by.get(k, 0) for k in ("high", "medium", "low")},
         "possibly real (separate list)": len(maybe), "saved": saved}, a.json)


def cmd_sample(a):
    *_, people = load(a.file, a)
    items = list(people.values())
    rnd = random.Random(a.seed)
    chosen = rnd.sample(items, min(a.n, len(items)))
    lines = [f"# Sample for manual review: {len(chosen)} of {len(items)} profiles", "",
             "Open each profile and fill in the columns. Type: personal / specialist / business / creator / bot?",
             "", "| # | Profile | Type | Field | City | Notes |", "|---|---|---|---|---|---|"]
    for i, p in enumerate(chosen, 1):
        link = p["url"] if p["url"].startswith("http") else f"https://www.instagram.com/{p['username']}/"
        label = (p["name"] or p["username"]).replace("|", "/")
        lines.append(f"| {i} | [{label}]({link}) |  |  |  |  |")
    text = "\n".join(lines) + "\n"
    if a.out:
        open(a.out, "w", encoding="utf-8").write(text)
        print(f"Saved: {a.out}")
    else:
        print(text)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--name-col"); p.add_argument("--url-col"); p.add_argument("--status-col")
    p.add_argument("--json", action="store_true")
    sub = p.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("inspect"); s.add_argument("file"); s.set_defaults(fn=cmd_inspect)
    s = sub.add_parser("changes"); s.add_argument("file"); s.add_argument("--out-dir")
    s.add_argument("--map", action="append", default=[]); s.add_argument("--show", type=int, default=30)
    s.set_defaults(fn=cmd_changes)
    s = sub.add_parser("compare"); s.add_argument("old"); s.add_argument("new")
    s.add_argument("--out-gained"); s.add_argument("--out-lost"); s.set_defaults(fn=cmd_compare)
    s = sub.add_parser("mutual"); s.add_argument("followers"); s.add_argument("following"); s.set_defaults(fn=cmd_mutual)
    s = sub.add_parser("profile"); s.add_argument("file"); s.add_argument("--examples", type=int, default=5)
    s.set_defaults(fn=cmd_profile)
    s = sub.add_parser("suspects"); s.add_argument("file"); s.add_argument("--out", required=True)
    s.add_argument("--lang", choices=["en", "ru"], default="en"); s.set_defaults(fn=cmd_suspects)
    s = sub.add_parser("sample"); s.add_argument("file"); s.add_argument("--n", type=int, default=40)
    s.add_argument("--seed", type=int, default=None); s.add_argument("--out"); s.set_defaults(fn=cmd_sample)
    a = p.parse_args(); a.fn(a)


if __name__ == "__main__":
    try:
        main()
    except BrokenPipeError:  # output was cut short (e.g. piped to head); not an error
        sys.stderr.close()
